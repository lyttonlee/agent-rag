"""导出下载路由：/v1/export。

导出 = 导出【已生成并存库】的内容：按 conversation_id + message_id 精确定位
某一轮起草的 assistant 消息正文，直接渲染 docx，**不重新检索/生成**。

安全：导出前对正文二次脱敏（敏感词表可能生成后更新）+ 所引文档密级复核，
不因内容已落库而放松出口管控。
"""

from __future__ import annotations

import io
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import SECURITY_WEIGHT
from app.core.security import redact_sensitive
from app.db.models import ConfigKV, Conversation, ConversationMessage, Document, Template
from app.db.session import get_db
from app.schemas.contracts import ExportDocxRequest
from app.services.auth import Caller, get_caller

router = APIRouter(prefix="/export", tags=["export"])


async def _active_sensitive_words(db: AsyncSession) -> list[str]:
    """解析当前生效的敏感词表（配置中心 category=sensitive_word 且 is_active）。"""
    rows = (
        await db.execute(
            select(ConfigKV.key).where(
                ConfigKV.category == "sensitive_word",
                ConfigKV.is_active.is_(True),
            )
        )
    ).scalars().all()
    return [w for w in rows if w]


async def _assert_template_usable(db: AsyncSession, template_code: str | None, caller: Caller) -> None:
    """导出版式模板所有权校验（他人模板不可用）。"""
    if not template_code:
        return
    row = (
        await db.execute(select(Template).where(Template.code == template_code))
    ).scalar_one_or_none()
    if row and not (
        row.owner_type == "system"
        or (
            row.owner_type == ("user" if caller.kind == "user" else "client")
            and row.owner_id == caller.id
        )
    ):
        raise HTTPException(status_code=403, detail="无权使用该模板（仅创建者与系统标准模板可用）")


def _clearance_allowed(clearance: str, caller_max: str) -> bool:
    """引用文档密级是否不超过调用方额度。"""
    return SECURITY_WEIGHT.get(clearance, 1) <= SECURITY_WEIGHT.get(caller_max, 1)


@router.post("/docx")
async def export_docx(
    req: ExportDocxRequest, caller: Caller = Depends(get_caller),
    db: AsyncSession = Depends(get_db),
) -> Response:
    """导出已生成内容（按会话+消息定位，不重新生成）为 docx。

    权限：第三方需具备 chat 能力域。安全：导出前对正文二次脱敏 + 所引文档
    密级复核，不因内容已落库而放松。
    """
    if not caller.allows_scope("chat"):
        raise HTTPException(status_code=403, detail="该客户端无 chat 能力域权限")
    await _assert_template_usable(db, req.template_code, caller)
    return await _run_export(req, caller, db)


async def _run_export(
    req: ExportDocxRequest, caller: Caller, db: AsyncSession
) -> Response:
    # 1) 定位会话
    conv = (
        await db.execute(select(Conversation).where(Conversation.id == req.conversation_id))
    ).scalar_one_or_none()
    if not conv:
        raise HTTPException(status_code=404, detail="会话不存在")

    # 2) 定位目标 assistant 消息（版本）：指定优先，否则取最新一条
    if req.message_id:
        target = (
            await db.execute(
                select(ConversationMessage).where(
                    ConversationMessage.id == req.message_id,
                    ConversationMessage.conversation_id == req.conversation_id,
                )
            )
        ).scalar_one_or_none()
        if not target or target.role != "assistant":
            raise HTTPException(status_code=400, detail="指定的消息不存在或非起草正文")
    else:
        target = (
            await db.execute(
                select(ConversationMessage)
                .where(
                    ConversationMessage.conversation_id == req.conversation_id,
                    ConversationMessage.role == "assistant",
                )
                .order_by(ConversationMessage.created_at.desc())
                .limit(1)
            )
        ).scalar_one_or_none()
    if not target or not (target.content or "").strip():
        raise HTTPException(status_code=400, detail="该会话尚无可导出的起草内容")

    # 3) 安全复核：二次脱敏 + 密级复核（防御纵深）
    words = await _active_sensitive_words(db)
    content, hits = redact_sensitive(target.content or "", words)
    doc_ids = [
        c.get("document_id") for c in (target.citations or []) if isinstance(c, dict) and c.get("document_id")
    ]
    if doc_ids:
        docs = (
            await db.execute(select(Document.id, Document.security_level).where(Document.id.in_(doc_ids)))
        ).all()
        for doc_id, sec in docs:
            if not _clearance_allowed(sec, caller.max_clearance):
                raise HTTPException(
                    status_code=403, detail="导出内容含超出您密级额度的引用，已拒绝导出"
                )

    # 4) 渲染 docx（docx 的 Document 改名以避免遮蔽 ORM 模型 Document）
    from docx import Document as DocxDocument

    doc = DocxDocument()
    title = conv.title or "公文起草"
    doc.add_heading(title, level=1)
    for para in (content or "").split("\n"):
        if para.strip():
            doc.add_paragraph(para)
    if target.citations:
        doc.add_heading("参考资料", level=2)
        for i, c in enumerate(target.citations, 1):
            if not isinstance(c, dict):
                continue
            doc.add_paragraph(f"[{i}] 《{c.get('title') or '未知'}》{c.get('doc_no') or ''}")
    if hits:
        doc.add_paragraph(f"【敏感词命中】{', '.join(hits)}（已脱敏）")

    # 直接渲染到内存字节流（避免临时文件落盘与并发句柄问题）
    safe_name = (title[:20] or "公文") + ".docx"
    buf = io.BytesIO()
    doc.save(buf)
    blob = buf.getvalue()
    # RFC 6266：filename 仅允许 ASCII，中文名放到 filename*（UTF-8 编码），
    # 避免 Starlette 头编码（Latin-1）失败导致 500
    disp = f"attachment; filename=\"document.docx\"; filename*=UTF-8''{quote(safe_name)}"
    return Response(
        content=blob,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": disp},
    )
