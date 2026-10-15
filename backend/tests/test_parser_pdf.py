"""解析器单测：PDF 位置感知清洗、DOCX 表格、TXT BOM、格式分发/嗅探。"""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from make_pdf import make_pdf

from app.services.parser import (
    _docx_table_lines,
    parse_bytes,
    parse_docx_bytes,
    parse_pdf_bytes,
    parse_text_bytes,
)
from ofdkit import looks_like_ofd, make_ofd

# 公文样例行：标题 + 正文（y 为基线坐标，页高 842）
BODY_LINES = [
    (72.0, 100.0, "四川省人民政府办公厅"),
    (72.0, 130.0, "关于加快提升农业社会化服务水平的指导意见"),
    (72.0, 170.0, "各市(州)人民政府，省政府各部门、各直属机构:"),
    (72.0, 200.0, "为加快健全我省便捷高效的农业社会化服务体系，经省政府同意，"),
    (72.0, 230.0, "现提出如下指导意见。"),
]


class TestPdfParsing:
    def test_basic_extraction(self):
        """基本提取：正文完整、页数正确。"""
        parsed = parse_pdf_bytes(make_pdf([BODY_LINES]))
        assert parsed.pages == 1
        assert "四川省人民政府办公厅" in parsed.text
        assert "现提出如下指导意见。" in parsed.text

    def test_page_numbers_removed(self):
        """页脚 `- N -` 页码剔除（多页）。"""
        pages = [[(72.0, 300.0, f"第{i}页正文内容，各市州人民政府有关部门。")] for i in range(1, 4)]
        parsed = parse_pdf_bytes(make_pdf(pages))
        assert "- 1 -" not in parsed.text
        assert "- 3 -" not in parsed.text
        assert "第1页正文内容" in parsed.text

    def test_repeated_header_removed(self):
        """每页重复的页眉行剔除。"""
        pages = [[(72.0, 300.0, f"第{i}页正文内容。")] for i in range(1, 4)]
        parsed = parse_pdf_bytes(make_pdf(pages, header="四川省人民政府办公厅文件"))
        assert "四川省人民政府办公厅文件" not in parsed.text
        assert parsed.text.count("正文内容") == 3

    def test_body_not_removed(self):
        """正文内容零误删（清洗字符损失率红线）。"""
        parsed = parse_pdf_bytes(make_pdf([BODY_LINES] * 3))
        for line in BODY_LINES:
            assert line[2] in parsed.text

    def test_noise_clean_disabled_keeps_noise(self):
        """关闭清洗：页码保留（对照审计）。"""
        parsed = parse_pdf_bytes(make_pdf([BODY_LINES]), noise_clean=False)
        assert "- 1 -" in parsed.text

    def test_scanned_detection(self):
        """空白页计为无文字层页。"""
        pages = [BODY_LINES, [(72.0, 100.0, "")]]
        parsed = parse_pdf_bytes(make_pdf(pages, page_number=False))
        assert parsed.pages == 2
        assert parsed.scanned_pages == 1

    def test_normalization_applied(self):
        """清洗顺带做全角半角规范化。"""
        page = [(72.0, 300.0, "川办发〔２０２５〕３６号文件已印发各市州。")]
        parsed = parse_pdf_bytes(make_pdf([page], page_number=False))
        assert "川办发〔2025〕36号" in parsed.text


class TestDocxParsing:
    def _make_docx(self) -> bytes:
        doc = DocxDocument()
        doc.add_paragraph("四川省人民政府办公厅")
        doc.add_paragraph("关于某事项的通知")
        doc.add_paragraph("各市(州)人民政府:")
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "指标"
        table.cell(0, 1).text = "目标值"
        table.cell(1, 0).text = "服务面积"
        table.cell(1, 1).text = "5000万亩次"
        table.cell(2, 0).text = "服务组织"
        table.cell(2, 1).text = "1000个"
        doc.add_paragraph("特此通知。")
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_table_extracted_in_order(self):
        """表格内容按文档顺序进入正文（doc.paragraphs 拿不到表格）。

        表格行格式：「表头1：值1；表头2：值2」——表头与值同行共生，
        检索时"服务面积/5000万亩次"共现，跨页重复表头去重。
        """
        parsed = parse_docx_bytes(self._make_docx())
        assert "四川省人民政府办公厅" in parsed.text
        assert "指标：服务面积；目标值：5000万亩次" in parsed.text
        assert "指标：服务组织；目标值：1000个" in parsed.text
        # 顺序：表格在"特此通知。"之前
        assert parsed.text.index("服务面积") < parsed.text.index("特此通知。")

    def test_docx_paragraphs_separated_by_blank_line(self):
        """DOCX 段落以空行连接——显式边界传给段落重建，防止误粘段。"""
        parsed = parse_docx_bytes(self._make_docx())
        assert "\n\n" in parsed.text

    def test_table_lines_dedup_header(self):
        """跨页重复表头行去重。"""
        from docx import Document as D

        doc = D()
        t = doc.add_table(rows=3, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "指标", "目标值"
        t.cell(1, 0).text, t.cell(1, 1).text = "服务面积", "5000"
        t.cell(2, 0).text, t.cell(2, 1).text = "指标", "目标值"  # 跨页重复表头
        lines = _docx_table_lines(t)
        assert lines == ["指标：服务面积；目标值：5000"]


class TestTextParsing:
    def test_bom_stripped(self):
        """UTF-8 BOM 剥离（否则首段带不可见字符）。"""
        parsed = parse_text_bytes("第一段。\n第二段。".encode("utf-8-sig"))
        assert parsed.text == "第一段。\n第二段。"

    def test_gbk_fallback(self):
        """GBK 编码兼容。"""
        parsed = parse_text_bytes("各市州人民政府。".encode("gbk"))
        assert parsed.text == "各市州人民政府。"

    def test_crlf_normalized(self):
        """CRLF 统一为 LF。"""
        parsed = parse_text_bytes(b"a\r\nb\rc")
        assert parsed.text == "a\nb\nc"


class TestDispatch:
    def test_ofd_dispatch(self):
        """.ofd 扩展名 → ofdkit 解析。"""
        ofd = make_ofd([[(20.0, 100.0, "四川省人民政府关于某事项的通知。")]])
        parsed = parse_bytes("doc.ofd", ofd)
        assert "四川省人民政府关于某事项的通知。" in parsed.text

    def test_ofd_sniff_without_extension(self):
        """无扩展名的 OFD（PK 魔数）→ 嗅探为 ofd 而非 docx。"""
        ofd = make_ofd([[(20.0, 100.0, "无扩展名公文内容测试段落。")]])
        parsed = parse_bytes("noext", ofd)
        assert "无扩展名公文内容测试段落。" in parsed.text
        assert looks_like_ofd(ofd) is True

    def test_docx_not_misdetected_as_ofd(self):
        """DOCX（同为 PK 魔数）不会被误判为 OFD。"""
        doc = DocxDocument()
        doc.add_paragraph("测试文档")
        buf = io.BytesIO()
        doc.save(buf)
        assert looks_like_ofd(buf.getvalue()) is False

    def test_pdf_sniff_without_extension(self):
        """无扩展名的 PDF → 嗅探为 pdf。"""
        parsed = parse_bytes("noext", make_pdf([BODY_LINES]))
        assert "四川省人民政府办公厅" in parsed.text

    def test_unsupported_format_raises(self):
        """不支持的格式明确报错。"""
        with pytest.raises(ValueError, match="暂不支持的解析格式"):
            parse_bytes("doc.wps", b"\x00\x01\x02random")
