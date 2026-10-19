"""上传参考文档（格式模板 / 仿写）——画像提取 + 路由 + 生成链路注入测试。

用户 2026-10-19 需求：
① 起草时上传一份文档，按该文档的**模板格式**生成新文档；
② **仿写**上传文档；③ 用户可以不选已有文种/模板。

不连真实 DB / 不调真实模型（与 test_draft_multiturn_integration 同口径）：
- 画像提取用 python-docx 现场构造参考文档（真实版式，非 mock）；
- 路由/生成链路用桩替身（存储、会话、模型、检索）。
"""

from __future__ import annotations

import io
from types import SimpleNamespace

import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.api.routes.generate import (
    get_reference,
    upload_reference,
)
from app.schemas.contracts import GenerateRequest


# =========================================================================
# 画像提取（纯函数）
# =========================================================================


def _set_cjk(run, name: str, size_pt: float, *, bold: bool = False,
             color: RGBColor | None = None) -> None:
    """给 run 设置中文字体（eastAsia 必写，否则 Word 渲染中文掉默认宋体）。"""
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def _make_reference_docx() -> bytes:
    """构造一份带版式的公文参考文档（红头 + 题名 + 黑体标题 + 仿宋正文）。"""
    import docx

    doc = docx.Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(3.7), Cm(3.5)
    sec.left_margin, sec.right_margin = Cm(2.8), Cm(2.6)

    # 红头机关标志（红色大字）
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cjk(p.add_run("XX集团有限公司文件"), "宋体", 26, bold=True,
             color=RGBColor(0xFF, 0x00, 0x00))
    # 题名（二号小标宋居中）
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_cjk(p.add_run("关于开展安全生产大检查的通知"), "方正小标宋简体", 22)
    # 一级标题（黑体）
    p = doc.add_paragraph(style="Heading 1")
    _set_cjk(p.add_run("一、检查范围"), "黑体", 16)
    # 正文（仿宋三号）
    for text in ("各分公司、各部门：", "为切实加强安全生产管理，现就开展安全"
                 "生产大检查有关事项通知如下。", "特此通知。"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _set_cjk(p.add_run(text), "仿宋_GB2312", 16)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_profile_extracts_layout():
    """DOCX 参考文档：版式画像（字体/红头/页边距）+ 结构大纲应被正确提取。"""
    from app.services.reference_profile import extract_reference_profile

    fmt, profile = extract_reference_profile("安全大检查通知.docx", _make_reference_docx())
    assert fmt == "docx"
    assert profile["has_real_layout"] is True
    layout = profile["layout"]
    assert layout["font"] == "仿宋_GB2312"           # 正文字体（众数）
    assert layout["title_font"] == "方正小标宋简体"    # 题名字体（红头后移一段）
    assert layout["red_head"] is True                # 红头检测
    assert layout["heading_font"] == "黑体"           # 层级标题字体
    # 结构大纲：黑体标题行应入选
    assert any("检查范围" in o["text"] for o in profile["outline"])
    # 规模与行文摘录
    assert profile["para_count"] >= 5
    assert profile["approx_words"] > 30
    assert "红头" in profile["style_notes"]
    assert "安全生产" in profile["excerpt"]


def test_docx_profile_keeps_gongwen_default_for_western_font():
    """正文字体为西文默认（Calibri）时不覆盖公文字体——否则中文掉回默认宋体。"""
    import docx

    from app.services.reference_profile import extract_reference_profile

    doc = docx.Document()
    p = doc.add_paragraph()
    run = p.add_run("纯西文字体的参考文档正文。")
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    buf = io.BytesIO()
    doc.save(buf)

    _fmt, profile = extract_reference_profile("plain.docx", buf.getvalue())
    assert "font" not in (profile["layout"] or {})
    assert profile["has_real_layout"] is False


def test_text_profile_outline_and_no_layout():
    """TXT 参考文档：无版式（None），但结构大纲/摘录仍可用于仿写。"""
    from app.services.reference_profile import extract_reference_profile

    text = (
        "XX单位关于开展消防检查的通知\n"
        "一、检查范围\n"
        "本次检查覆盖全部办公区域与仓库。\n"
        "二、整改要求\n"
        "各单位须在十个工作日内完成整改并书面反馈。\n"
        "特此通知。\n"
    ).encode("utf-8")
    fmt, profile = extract_reference_profile("参考.txt", text)
    assert fmt == "text"
    assert profile["layout"] is None
    assert profile["has_real_layout"] is False
    assert any("检查范围" in o["text"] for o in profile["outline"])
    assert any("整改要求" in o["text"] for o in profile["outline"])
    assert "特此通知" in profile["closing"]


@pytest.mark.parametrize("name,data,err_part", [
    ("broken.docx", b"PK\x03\x04garbage-not-a-real-zip", "解析失败"),
    ("bad.txt", b"", "文档为空"),
])
def test_invalid_reference_raises(name, data, err_part):
    """空文件/损坏文件必须显式报错（上传即得反馈，不留中间态）。"""
    from app.services.reference_profile import extract_reference_profile

    with pytest.raises(ValueError) as exc:
        extract_reference_profile(name, data)
    assert err_part in str(exc.value)


# =========================================================================
# 提示词指令渲染
# =========================================================================


def _fake_profile(**over):
    base = {
        "layout": {"font": "仿宋_GB2312", "title_font": "方正小标宋简体", "red_head": True},
        "outline": [{"level": 2, "text": "一、检查范围"}],
        "opening": "各分公司、各部门：",
        "closing": "特此通知。",
        "excerpt": "为切实加强安全生产管理……",
        "style_notes": "标题：方正小标宋简体 22pt（居中）；正文：仿宋_GB2312 16pt",
        "para_count": 5,
        "approx_words": 100,
        "has_real_layout": True,
    }
    base.update(over)
    return base


def test_render_instructions_modes():
    """三种用法渲染出不同指令块；format 必须带版式、imitate 必须带摘录。"""
    from app.services.reference_profile import render_reference_instructions

    fmt_text = render_reference_instructions(_fake_profile(), "format")
    assert "【版式要求】" in fmt_text and "仿宋_GB2312" in fmt_text
    assert "【仿写要求】" not in fmt_text

    imi_text = render_reference_instructions(_fake_profile(), "imitate")
    assert "【仿写要求】" in imi_text
    assert "为切实加强安全生产管理" in imi_text      # 代表性段落摘录
    assert "【版式要求】" not in imi_text             # 纯仿写不强制版式

    both = render_reference_instructions(_fake_profile(), "both")
    assert "【版式要求】" in both and "【仿写要求】" in both


def test_render_instructions_format_without_layout_not_silent():
    """PDF/TXT 参考选 format 模式：不能静默无指令，需明示按公文标准版式。"""
    from app.services.reference_profile import render_reference_instructions

    text = render_reference_instructions(_fake_profile(has_real_layout=False), "format")
    assert "【版式说明】" in text and "GB/T 9704" in text


# =========================================================================
# 路由（桩替身：存储 + DB）
# =========================================================================


class _FakeStorage:
    def __init__(self):
        self.saved: dict[str, bytes] = {}

    def save(self, key: str, data: bytes) -> str:
        self.saved[key] = data
        return key

    def get(self, key):  # pragma: no cover - 路由未用
        return self.saved[key]


class _FakeDB:
    """最小 DB 桩：add/commit/refresh + execute(select)。"""

    def __init__(self):
        self.rows: list = []

    def add(self, row):
        self.rows.append(row)

    async def commit(self):
        pass

    async def refresh(self, row):
        pass

    async def execute(self, stmt):
        return SimpleNamespace(scalar_one_or_none=lambda: self.rows[0] if self.rows else None)


class _File:
    def __init__(self, filename: str, data: bytes):
        self.filename = filename
        self._data = data

    async def read(self) -> bytes:
        return self._data


def _caller(kind: str = "user", cid: str = "u-1"):
    return SimpleNamespace(kind=kind, id=cid, name="测试用户", allows_scope=lambda s: True)


@pytest.mark.asyncio
async def test_upload_reference_route(monkeypatch):
    """上传 DOCX 参考：落库 + 画像 + mode_default=format（有版式）。"""
    from app.services import reference_profile  # noqa: F401  (确保模块已加载)

    storage = _FakeStorage()
    monkeypatch.setattr("app.api.routes.generate.get_storage", lambda: storage)
    db = _FakeDB()

    resp = await upload_reference(
        file=_File("安全大检查通知.docx", _make_reference_docx()),
        caller=_caller(),
        db=db,
    )
    assert resp.status == "PARSED"
    assert resp.format == "docx"
    assert resp.mode_default == "format"
    assert resp.profile["has_real_layout"] is True
    assert resp.profile["style_notes"]
    # 原文进存储、行进 draft_references
    assert len(storage.saved) == 1 and len(db.rows) == 1
    row = db.rows[0]
    assert row.owner_type == "user" and row.owner_id == "u-1"
    assert row.storage_key in storage.saved


@pytest.mark.asyncio
async def test_upload_reference_rejects_unsupported(monkeypatch):
    """垃圾字节（嗅探不出格式）→ 400，不落库不落存储。"""
    from fastapi import HTTPException

    storage = _FakeStorage()
    monkeypatch.setattr("app.api.routes.generate.get_storage", lambda: storage)
    db = _FakeDB()

    with pytest.raises(HTTPException) as exc:
        await upload_reference(
            file=_File("virus.bin", b"MZ\x90\x00 not a doc"),
            caller=_caller(),
            db=db,
        )
    assert exc.value.status_code == 400
    assert not storage.saved and not db.rows


@pytest.mark.asyncio
async def test_get_reference_owner_isolation(monkeypatch):
    """他人参考文档 → 404（不可见=不存在，与空间/模板口径一致）。"""
    from fastapi import HTTPException

    from app.db.models import DraftReference

    row = DraftReference(
        id="ref-1", filename="a.docx", format="docx", size_bytes=10,
        storage_key="references/ref-1.docx", profile={"has_real_layout": True},
        status="PARSED", owner_type="user", owner_id="someone-else",
    )

    class _DB(_FakeDB):
        async def execute(self, stmt):
            return SimpleNamespace(scalar_one_or_none=lambda: row)

    with pytest.raises(HTTPException) as exc:
        await get_reference(reference_id="ref-1", caller=_caller("user", "me"), db=_DB())
    assert exc.value.status_code == 404


# =========================================================================
# 生成链路注入
# =========================================================================


@pytest.mark.asyncio
async def test_generate_with_reference_injects_instructions(monkeypatch):
    """单趟链路：reference_id → 成稿 messages 必须含【参考文档要求】块。

    单趟（deep_reason 缺省时跟配置，测试环境无配置库→按默认 True 会走
    ReAct，故显式 deep_reason=false 走单趟桩链路，聚焦验证提示词注入）。
    """
    from app.db.models import DraftReference
    from app.services import generation

    store: list[tuple[str, str, str]] = []
    profile = _fake_profile()

    async def fake_append(db, cid, role, content, **kw):
        store.append((cid, role, content))

    async def fake_recent(db, cid, limit=10):
        return []

    async def fake_create(*a, **kw):
        return SimpleNamespace(id="conv-A")

    async def _ensure(db, req, doc_type=None):
        return "conv-A"

    async def _resolve_words(db, req):
        return []

    async def _agent_off(db):
        return SimpleNamespace(enabled=False)

    async def _retrieve(db, req):
        return SimpleNamespace(citations=[], truncated_by_security=False)

    async def _load_react_cfg(db):
        from app.services.agent.react_config import ReactConfig

        return ReactConfig(draft_enabled=False)  # 单趟链路

    monkeypatch.setattr(generation.conv_service, "create_conversation", fake_create)
    monkeypatch.setattr(generation.conv_service, "append_message", fake_append)
    monkeypatch.setattr(generation.conv_service, "recent_history", fake_recent)
    monkeypatch.setattr(generation, "_ensure_draft_conversation", _ensure)
    monkeypatch.setattr(generation, "_resolve_sensitive_words", _resolve_words)
    monkeypatch.setattr(generation, "load_agent_config", _agent_off)
    monkeypatch.setattr(generation, "retrieve_enhanced", _retrieve)
    monkeypatch.setattr(generation, "load_react_config", _load_react_cfg)

    captured: dict = {}

    class _Gen:
        model_name = "mock-gen"

        async def generate_stream(self, messages, **kw):
            captured["messages"] = messages
            captured["kwargs"] = kw
            yield ("content", "一、测试")

    monkeypatch.setattr(generation, "ModelProvider", lambda role: _Gen())

    # 参考文档行：_load_draft_reference 走 db.execute
    row = DraftReference(
        id="ref-A", filename="参考.docx", format="docx", size_bytes=10,
        storage_key="references/ref-A.docx", profile=profile,
        status="PARSED", owner_type="user", owner_id="u-1",
    )

    async def fake_execute(stmt):
        return SimpleNamespace(scalar_one_or_none=lambda: row)

    class _DB:
        async def execute(self, stmt):
            return await fake_execute(stmt)

    req = GenerateRequest(
        query="起草一份安全生产检查通知", reference_id="ref-A", deep_reason=False
    )
    chunks = []
    async for chunk in generation.stream_generate(_DB(), req, actor="t", caller=_caller()):
        chunks.append(chunk)

    assert captured.get("messages"), "模型未被调用"
    system = captured["messages"][0]["content"]
    assert "【参考文档要求】" in system
    assert "【版式要求】" in system          # mode 默认 format（有版式）
    assert "仿宋_GB2312" in system           # 参考文档正文字体进入提示词


@pytest.mark.asyncio
async def test_generate_reference_ownership_404(monkeypatch):
    """引用他人 reference_id → 404，且不进入 SSE 流（响应头未发）。"""
    from fastapi import HTTPException

    from app.db.models import DraftReference
    from app.services import generation

    row = DraftReference(
        id="ref-B", filename="他人.docx", format="docx", size_bytes=1,
        storage_key="references/ref-B.docx", profile={},
        status="PARSED", owner_type="user", owner_id="someone-else",
    )

    async def fake_execute(stmt):
        return SimpleNamespace(scalar_one_or_none=lambda: row)

    class _DB:
        async def execute(self, stmt):
            return await fake_execute(stmt)

        async def commit(self):
            pass

    req = GenerateRequest(query="随便", reference_id="ref-B")
    with pytest.raises(HTTPException) as exc:
        async for _ in generation.stream_generate(_DB(), req, actor="t", caller=_caller()):
            pass
    assert exc.value.status_code == 404


@pytest.mark.asyncio
async def test_react_draft_task_context_includes_reference_note():
    """ReAct 任务上下文：有参考文档时必须带【格式/风格来源】说明。"""
    from app.schemas.contracts import GenerateRequest
    from app.services.generation import _build_draft_task_context

    req = GenerateRequest(query="起草通知")
    note = "【格式/风格来源】用户上传了参考文档《模板.docx》（按该文档的模板格式生成）"
    ctx = _build_draft_task_context(req, None, False, reference_note=note)
    assert note in ctx
    # 无参考文档时行为不变
    plain = _build_draft_task_context(req, None, False)
    assert "【格式/风格来源】" not in plain
